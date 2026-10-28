from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path
root=Path(__file__).resolve().parents[1]
d=Document();section=d.sections[0];section.top_margin=Inches(.72);section.bottom_margin=Inches(.7);section.left_margin=Inches(.85);section.right_margin=Inches(.85)
styles=d.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10)
for nm,size in [('Title',28),('Heading 1',16),('Heading 2',12)]:
 st=styles[nm];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor(18,67,85)
def p(t='',style=None):return d.add_paragraph(t,style)
def h(t,l=1):d.add_heading(t,l)
def table(headers,rows):
 t=d.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,x in zip(t.rows[0].cells,headers):c.text=x
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=str(x)
 d.add_paragraph('')
d.add_heading('RetailIQ Copilot',0)
p('Domain Selection · Business Problem · Architecture · Tech Stack · Week 1 Demo')
p('AI Flezi | Domain #15 Retail / E-commerce | Electronics & Accessories | October 2026')
p('Purpose: concise working document for a team of five with 60 minutes per person/day, including course learning. Local demo only.')
h('1. Decision and alignment')
table(['Item','Decision'],[['Official project','RetailIQ Copilot – Omnichannel Inventory Intelligence & Conversational Analytics'],['Domain','#15 Retail / E-commerce'],['Product vertical','Electronics & Accessories'],['Business model','Online channel + HCMC and Hanoi stores + central warehouse'],['Week 1 outcome','Inventory & Sales Assistant with Natural Language Queries'],['Deferred prior idea','Order/Delivery Exception Assistance: later Order Management and Fulfilment scope']])
h('2. Business problem (hypothesis)')
p('Inventory and sales stakeholders must reconcile multiple systems to answer common questions about available stock, monthly revenue, top sellers, and technical product details. Manual aggregation creates potential delay, inconsistent metric definitions, freshness uncertainty and fragmented knowledge. This is a plausible business problem for a synthetic retailer, not a validated finding from a real enterprise.')
h('3. Stakeholder / AS-IS / TO-BE')
table(['Stakeholder','Primary question'],[['Inventory manager','How much stock is available by SKU and location?'],['Sales manager','What is completed-order net revenue this month?'],['Store manager','Where is product stock located?'],['Merchandising','What are product specifications and policy notes?']])
p('AS-IS (assumed): Ask → inspect stock report / sales report / product files → filter by date/location → reconcile KPI definitions → manually reply.')
p('TO-BE Week 1: Ask in natural language → LangGraph chooses bounded tool → read-only SQLite SQL or Chroma product RAG → source-grounded response. RAGAS evaluates product answer quality; LangSmith provides execution traces when configured.')
h('4. Business KPI contracts')
table(['Metric','Definition / caveat'],[['Available units','on_hand – reserved, group by SKU/location, show snapshot timestamp'],['Net sales revenue','SUM(quantity × unit_price – line_discount) on completed orders; VND; excludes tax, shipping and returns'],['Top sellers','SUM(quantity) by SKU, completed orders, chosen calendar month'],['Product knowledge','Use retrieved technical docs; never infer live inventory or sales from product documents']])
d.add_page_break()
h('5. Week 1 reference architecture')
p('Local browser → FastAPI → LangGraph StateGraph (route) → SQL Tool (structured inventory/sales) OR Product RAG (LangChain + Chroma VectorDB + embeddings) → grounded response. LangSmith traces graph execution; RAGAS separately scores grounded product answers.')
h('6. Tech stack')
table(['Component','Technology','Reason'],[['Runtime','Python 3.11+','Agent ecosystem and rapid local setup'],['Orchestration','LangGraph 1.x StateGraph','Explicit SQL/RAG conditional graph'],['Tools','LangChain @tool','Bounded read-only SQL analytics'],['Database','SQLite','Portable reproducible dataset'],['Retrieval','LangChain + Chroma','Local vector store of product knowledge'],['Model/embeddings','LangChain OpenAI','LLM intent + grounded product response'],['Tracing','LangSmith','Hierarchical run/tool tracing'],['Evaluation','RAGAS','Product RAG faithfulness + answer relevancy'],['UI','FastAPI + static HTML/JS','Local browser, minimal dependencies'],['Tests','pytest','Business logic and data correctness']])
h('7. Dataset design')
p('Synthetic seed: 10 products (laptops, phones, peripherals, accessories), 3 stock locations, 30 SKU–location inventory rows, 120 generated orders, completed/cancelled statuses, online/store sales channels, and product markdown documents. Data as-of is fixed to 2026-10-28 to keep the demo repeatable.')
d.add_page_break()
h('8. Weekly extension contract')
table(['Week','Official capabilities'],[['1','LangGraph + SQL, product Vector RAG, RAGAS, LangSmith'],['2','Freshness tiers, semantic layer, MCP Inventory/Order, context memory, hybrid retrieval, rewrite/rerank, accuracy report'],['3','HITL pricing, row-level security, KPI alerts, cost/token optimization, LangFuse, DeepEval, governance/security reports'],['4','Five agents: Inventory Search, Fulfilment Optimisation, Customer Context, Dynamic Pricing, Order Management; A2A, verification, report and manager approval']])
h('9. Demo script — 6 minutes')
table(['Time','Demo','Evidence'],[['0:00–1:20','Domain + problem + architecture','One local project with defined KPI contracts'],['1:20–2:10','Stock for AlphaBook Pro 14','Available total, breakdown by location'],['2:10–3:00','Revenue this month','Completed orders / by-channel breakdown'],['3:00–3:45','Top sellers','Rank by units sold'],['3:45–4:30','USB-C GaN Charger product info','RAG retrieved document with source'],['4:30–5:20','Testing and evaluation','pytest; real RAGAS & LangSmith only if configured'],['5:20–6:00','Limitations and roadmap','Weeks 2–4 build on same data and contracts']])
h('10. Verification status at packaging')
p('PASS: local synthetic SQLite seed, HTTP UI/API, four query examples and six deterministic pytest tests. NOT YET VERIFIED: actual LangGraph + Chroma + RAGAS + LangSmith running together with real model/tracing credentials, because those packages/credentials are unavailable in this build environment. The offline fallback must never be presented as evidence of those rubric features.')
h('11. Next validation checkpoint')
p('Before the Week 1 demo, install full dependencies on a connected workstation, configure model and LangSmith credentials, run every query in full mode, capture one LangSmith trace, run RAGAS and save score report. If this fails, disclose the current state rather than presenting mock results.')
out=root/'docs'/'RetailIQ_Business_Architecture_Week1.docx';d.save(out);print(out)
