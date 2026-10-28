"""Rebuild Week 1 PPTX and insert official checklist coverage slides."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import re

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'docs/slides/RetailIQ_Flezi_Week1_Demo.pptx'
OUT=ROOT/'docs/slides/RetailIQ_Flezi_Week1_Checklist_Aligned.pptx'
prs=Presentation(BASE)
BG='0C1424'; CARD='17253B'; DARK='21344D'; TEXT='EAF2FF'; MUTE='9CADC6'; TEAL='72DED1'; GOLD='F4BF77'; LINE='30425B'

def block(sl,x,y,w,h,color,round=True):
    shp=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid();shp.fill.fore_color.rgb=RGBColor.from_string(color)
    shp.line.fill.background()
    return shp

def t(sl,content,x,y,w,h,size=18,color=TEXT,bold=False):
    shp=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));f=shp.text_frame
    f.word_wrap=True;f.margin_left=Inches(.06);f.margin_top=Inches(.04)
    for i,line in enumerate(content.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(5)
    return shp

def base(name,subtitle):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    block(s,.48,.45,.08,.43,TEAL,False)
    t(s,'RETAILIQ COPILOT  ·  FLEZI WEEK 1',.75,.42,8,.35,11,TEAL,True)
    t(s,name,.68,1.08,12,.67,30,TEXT,True)
    t(s,subtitle,.70,1.87,11.8,.55,14,MUTE)
    block(s,.66,7.14,12,.012,LINE,False)
    t(s,'DOMAIN #15  ·  RETAIL / E-COMMERCE',.72,7.16,8,.24,9,MUTE)
    t(s,'',11.75,7.16,.8,.24,9,MUTE)
    return s

def panel(s,x,y,w,h,title,body,accent=TEAL,title_size=17,body_size=14):
    block(s,x,y,w,h,CARD)
    block(s,x+.19,y+.24,.07,.25,accent,False)
    t(s,title,x+.37,y+.17,w-.55,.45,title_size,TEXT,True)
    t(s,body,x+.22,y+.72,w-.42,h-.81,body_size,MUTE)

# Created at indices 9..13 before ordering.
s=base('FDE: discover → deploy → learn','Problem-first delivery, rooted in a retailer’s daily decision workflow')
steps=[('DOMAIN','Electronics & Accessories\nOnline + physical-store inventory'),('DISCOVERY','Manual reporting hypothesis\nStakeholders and KPI definitions'),('FIELD PROTOTYPE','SQL analytics + Product RAG\n4 customer-facing queries'),('FEEDBACK','Test accuracy, review trace\nRefine failure cases')]
for i,(h,b) in enumerate(steps):
    panel(s,.74+i*3.17,2.64,2.98,2.73,h,b,TEAL if i in (0,2) else GOLD,14,13)
block(s,.81,5.79,11.66,.77,DARK)
t(s,'EVIDENCE LIMIT: business scenario and observations are synthetic; no real customer interviews claimed.',1.01,5.96,11.1,.45,14,GOLD,True)

s=base('Harness architecture: 5 practical layers','Course concepts mapped to concrete RetailIQ modules and trust boundaries')
rows=[('1  EXPERIENCE','HTML Chat UI + FastAPI','Prompt, request validation, response'),('2  CONTEXT','Graph State + KPI contracts + RAG chunks','Only relevant data to each path'),('3  ORCHESTRATION','LangGraph Router + conditional edges','Explicit flow and state'),('4  TOOLS & DATA','Read-only SQL / Chroma Product RAG','Authoritative structured facts'),('5  VERIFICATION & OPS','pytest / RAGAS / LangSmith / Cost','Testing, tracing and quality gates')]
for i,(h,detail,description) in enumerate(rows):
    y=2.41+i*.81
    block(s,.84,y,11.63,.7,CARD)
    t(s,h,1.05,y+.11,3.05,.38,14,TEAL if i%2==0 else GOLD,True)
    t(s,detail,4.03,y+.1,4.13,.45,15,TEXT,True)
    t(s,description,8.14,y+.13,3.9,.42,12,MUTE)

s=base('LangChain + LangGraph pattern decisions','Demonstrate concept selection without pretending Week 1 is multi-agent')
patterns=[('CHAIN','Retrieval → context → answer','Conceptual sequence'),('ROUTER','Stock / revenue / top / product','Implemented'),('CUSTOM WORKFLOW','LangGraph conditional graph','Implemented'),('TOOL-USING FLOW','LangChain @tool for SQL','Implemented'),('SUBAGENT / HANDOFF','Separate domain agents / transfers','Week 4'),('SKILL','Load specialist guidance on demand','Not implemented')]
for i,(name,desc,status) in enumerate(patterns):
    col=i%2; row=i//2;x=.79+col*6.2;y=2.45+row*1.35
    block(s,x,y,5.97,1.17,CARD)
    t(s,name,x+.23,y+.15,2.4,.35,14,TEAL if status=='Implemented' else GOLD,True)
    t(s,status.upper(),x+4.1,y+.15,1.64,.34,10,MUTE,True)
    t(s,desc,x+.2,y+.57,5.52,.39,13,TEXT)

s=base('Agent cost & ROI: assumptions matter','A worked financial model is required; no actual savings have been measured')
panel(s,.75,2.49,3.72,2.64,'USAGE COST','Monthly LLM: input tokens + output tokens; embeddings + evaluation + hosting/ops.',TEAL,16,14)
panel(s,4.65,2.49,3.72,2.64,'ILLUSTRATIVE SAVINGS','1,000 queries × 60% benefited × 4 minutes saved × 150k VND/hour = 6m VND.',GOLD,16,13)
panel(s,8.55,2.49,3.72,2.64,'PAYBACK','Net benefit = 6m − 1.2m = 4.8m VND/month. Build cost 18m → 3.75 months.',TEAL,16,13)
block(s,.8,5.54,11.6,.99,DARK)
t(s,'ALL VALUES ARE ASSUMPTIONS. Actual ROI needs LangSmith tokens, current model rates, measured handling time and human QA.',1.08,5.71,11.05,.68,14,GOLD,True)

s=base('Testing + limitations + next verification','What passed locally versus what still needs real provider and LangSmith evidence')
panel(s,.74,2.5,5.92,2.65,'VERIFIED LOCALLY','43/43 pytest tests: SQL math, URL/key isolation, API root + 4 question flows, 8 document loaders, ROI calculations.',TEAL,17,14)
panel(s,6.78,2.5,5.77,2.65,'REMAINING E2E GATES','Full LangGraph + Chroma with real models; actual RAGAS scores; LangSmith trace with separate key.',GOLD,17,14)
block(s,.79,5.49,11.7,.99,DARK)
t(s,'RISKS: synthetic and snapshot data, ambiguous intent, RAG hallucination, no production authentication. W2/W3 address freshness and governance.',1.08,5.7,11.08,.67,13,GOLD,True)

# Move inserted slides among the original 9 for coherent talk track.
order=[0,1,9,2,10,3,11,4,5,6,12,13,7,8]
ids=prs.slides._sldIdLst
children=list(ids)
assert len(children)==len(order)==14
for child in children:ids.remove(child)
for ix in order:ids.append(children[ix])

for slide_index,slide in enumerate(prs.slides,1):
    for sh in slide.shapes:
        if not sh.has_text_frame:continue
        txt0=sh.text
        if not txt0:continue
        new=txt0
        new=re.sub(r'^\d{2}  /  RETAILIQ COPILOT',f'{slide_index:02d}  /  RETAILIQ COPILOT',new)
        new=re.sub(r'\d{2} / 09$',f'{slide_index:02d} / 14',new)
        new=new.replace('VERIFIED NOW: 6 OFFLINE TESTS PASS.', 'VERIFIED NOW: 43 LOCAL TESTS PASS.')
        new=new.replace('6 deterministic tests passing','43 deterministic + format tests pass')
        if new!=txt0:
            # Keep original styling by changing runs in existing single paragraph when possible.
            if len(sh.text_frame.paragraphs)==1 and len(sh.text_frame.paragraphs[0].runs)==1:
                sh.text_frame.paragraphs[0].runs[0].text=new
            else:
                sh.text=new
    # Ensure new slides have a footer counter.
    if slide_index in [3,5,7,11,12]:
        t(slide,f'{slide_index:02d} / 14',11.75,7.16,.8,.24,9,MUTE)

OUT.parent.mkdir(exist_ok=True,parents=True)
prs.save(OUT)
print('Slides:',len(prs.slides),'Output:',OUT)
