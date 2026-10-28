from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
R=Presentation();R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
BG='#0C1424';CARD='#17253B';DARK='#21344D';TEXT='#EAF2FF';MUTE='#9CADC6';TEAL='#72DED1';GOLD='#F4BF77';LINE='#30425B'
def rgb(c):return RGBColor.from_string(c.strip('#'))
def shape(s,x,y,w,h,fill,rad=True,line=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=rgb(line)
 return sh
def txt(s,text,x,y,w,h,size=20,color=TEXT,bold=False,align=None):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=b.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.06);tf.margin_top=Inches(.04)
 for i,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(6)
  if align:p.alignment=align
 return b
def base(k,title,subtitle=''):
 s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 shape(s,.48,.45,.08,.43,TEAL,False)
 txt(s,f'{k:02d}  /  RETAILIQ COPILOT  ·  FLEZI WEEK 1',.75,.42,9,.35,11,TEAL,True)
 txt(s,title,.68,1.08,12,.67,31,TEXT,True)
 if subtitle:txt(s,subtitle,.7,1.87,11.9,.54,15,MUTE)
 shape(s,.66,7.14,12,.012,LINE,False)
 txt(s,'DOMAIN #15  ·  RETAIL / E-COMMERCE',.72,7.16,8,.24,9,MUTE)
 txt(s,f'{k:02d} / 09',11.75,7.16,.8,.24,9,MUTE)
 return s
def card(s,x,y,w,h,title,body,accent=TEAL):
 shape(s,x,y,w,h,CARD)
 shape(s,x+.2,y+.25,.07,.26,accent,False)
 txt(s,title,x+.39,y+.18,w-.6,.43,18,TEXT,True)
 txt(s,body,x+.25,y+.77,w-.48,h-.9,14,MUTE)
# 1 cover
s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
shape(s,.72,.8,.09,.55,TEAL,False);txt(s,'AI FLEZI  •  DOMAIN #15  •  WEEK 01',.98,.84,10,.35,14,TEAL,True)
txt(s,'RetailIQ',.9,1.74,11.5,1.05,54,TEXT,True);txt(s,'Copilot',.9,2.78,9.3,1.02,54,TEAL,True)
txt(s,'Omnichannel Inventory Intelligence\n& Conversational Analytics',.95,4.05,10.6,1.15,27,TEXT)
shape(s,.95,6.1,7.8,.61,CARD);txt(s,'LOCAL DEMO   /   SQL ANALYTICS   /   PRODUCT RAG   /   LANGGRAPH',1.15,6.25,7.6,.31,12,MUTE,True)
# 2
s=base(2,'Why Retail / E-commerce?','Domain selection  ·  Electronics & Accessories  ·  Omnichannel')
card(s,.75,2.64,3.82,2.65,'BUSINESS SETTING','Online storefront + two physical stores + one warehouse. Products have many SKUs and technical specifications.')
card(s,4.73,2.64,3.82,2.65,'STAKEHOLDERS','Inventory planners, sales managers, store operators and merchandising specialists.')
card(s,8.72,2.64,3.82,2.65,'VALUE HYPOTHESIS','Reduce manual report joins and ambiguous KPI answers. Business impact is not yet validated.',GOLD)
shape(s,.75,5.78,11.8,.65,DARK);txt(s,'NOTE  Synthetic retailer / synthetic sales data — not a claim about a real business.',1,5.94,11.2,.35,14,GOLD)
# 3
s=base(3,'Problem-first → measurable scope','What problem are we solving in Week 1?')
card(s,.75,2.58,5.7,3.27,'AS-IS / MANUAL','Analyst searches inventory sheets → opens sales reports → joins channel data → checks product descriptions → explains results.',GOLD)
card(s,6.64,2.58,5.85,3.27,'TO-BE / RETAILIQ','Ask in natural language → route to SQL or Product RAG → calculate governed metrics → show answer and evidence.',TEAL)
shape(s,.8,6.1,11.62,.52,DARK);txt(s,'SCOPE GUARD: READ-ONLY. NO PRICE CHANGES, ORDER EXECUTION OR AUTONOMOUS DECISIONS IN W1.',1.05,6.22,11.1,.3,12,TEXT,True)
# 4 architecture
s=base(4,'Week 1 architecture','Structured facts via vetted SQL; unstructured product knowledge via vector retrieval')
shape(s,.8,2.4,2,.74,DARK);txt(s,'Local Chat UI',.98,2.56,1.65,.31,16,TEXT,True)
shape(s,3.3,2.4,2.15,.74,DARK);txt(s,'FastAPI',3.63,2.56,1.6,.31,16,TEXT,True)
shape(s,5.96,2.4,2.22,.74,'#204A55');txt(s,'LangGraph',6.23,2.55,1.73,.31,16,TEXT,True)
shape(s,8.8,2.03,3.55,1.22,CARD);txt(s,'SQL Analytics Tool',9.12,2.15,3.1,.4,17,TEAL,True);txt(s,'SQLite / Inventory + Sales',9.1,2.69,3.1,.34,13,MUTE)
shape(s,8.8,3.58,3.55,1.2,CARD);txt(s,'Product RAG',9.12,3.69,3.0,.4,17,GOLD,True);txt(s,'Chroma + Product Docs',9.12,4.2,3.1,.32,13,MUTE)
for x,y,t in [(2.85,2.48,'→'),(5.52,2.48,'→'),(8.33,2.48,'→'),(8.33,3.58,'→')]:txt(s,t,x,y,.5,.45,24,TEAL,True)
shape(s,.81,5.33,5.53,.8,CARD);txt(s,'LangSmith  ·  tracing (enabled via API key)',1.06,5.55,5.0,.37,14,TEXT)
shape(s,6.63,5.33,5.73,.8,CARD);txt(s,'RAGAS  ·  faithfulness + relevancy',6.91,5.55,5.12,.37,14,TEXT)
txt(s,'Full mode needs dependencies + model credentials. Offline smoke demo is not a substitute.',.89,6.49,11.4,.38,12,GOLD)
# 5
s=base(5,'Data contracts & source of truth','Omnichannel means both store-level inventory and online + physical-store sales')
card(s,.76,2.52,3.8,2.95,'INVENTORY','Available = on hand − reserved.\nSKU × location.\nShow timestamp and per-location breakdown.')
card(s,4.75,2.52,3.8,2.95,'NET REVENUE','Completed orders only.\nSum(qty × price − discounts).\nVND; before returns, taxes and shipping.')
card(s,8.74,2.52,3.8,2.95,'TOP SELLERS','Rank by units sold, not by revenue.\nFilter to completed orders.\nCurrent configured calendar month.')
txt(s,'DEMO SNAPSHOT: 2026-10-28  ·  10 PRODUCTS  ·  3 LOCATIONS  ·  120 ORDERS',.84,6.01,11.6,.42,16,TEAL,True)
# 6
s=base(6,'Four queries, one assistant','Live demo: 4 tasks mapped to the exact Week 1 project outcome')
questions=[('01','INVENTORY','How many units of AlphaBook Pro 14 are currently in stock?'),('02','SALES','What is the revenue this month?'),('03','RANKING','Which products are top sellers?'),('04','PRODUCT RAG','Tell me about USB-C GaN Charger 65W')]
for i,(n,k,q) in enumerate(questions):
 y=2.42+i*1.02;shape(s,.84,y,11.55,.84,CARD);txt(s,n,1.04,y+.19,.48,.3,14,TEAL,True);txt(s,k,1.67,y+.18,2.0,.35,13,MUTE,True);txt(s,q,3.8,y+.15,8.3,.52,16,TEXT)
# 7
s=base(7,'Evidence, evaluation & limitations','Three distinct checks — do not mix deterministic correctness with LLM quality')
card(s,.74,2.57,3.84,2.75,'SQL & BUSINESS TESTS','Deterministic pytest checks stock math, completed-only revenue, ranking, intent and read-only SQL.')
card(s,4.74,2.57,3.84,2.75,'RAGAS','Faithfulness and answer relevancy over product questions. Requires provider key and full-mode run.')
card(s,8.75,2.57,3.84,2.75,'LANGSMITH','Record graph runs, prompt/tool calls and execution time. Requires credentials and tracing enabled.')
shape(s,.81,5.8,11.68,.74,DARK);txt(s,'VERIFIED NOW: 6 OFFLINE TESTS PASS.  FULL-MODE INTEGRATION STILL NEEDS LOCAL CREDENTIAL TEST.',1.03,5.99,11.23,.42,14,GOLD,True)
# 8
s=base(8,'Four-week system evolution','One codebase · mandatory capabilities added progressively')
for i,(head,body) in enumerate([('W1  AI ASSISTANT','LangGraph + SQL + Product RAG\nRAGAS · LangSmith'),('W2  CONVERSATIONAL ANALYTICS','Freshness tiers · MCP inventory/order\nSemantic layer · Memory · Hybrid retrieval'),('W3  GOVERNED COPILOT','HITL pricing · Row security · KPI alerts\nLangFuse · DeepEval · AgentOps'),('W4  RETAIL OPERATING SYSTEM','5 agents · A2A · Verification loop\nOrder E2E + manager approval')]):
 x=.76+(i%2)*6.1;y=2.46+(i//2)*2.0;card(s,x,y,5.94,1.8,head,body,TEAL if i%2==0 else GOLD)
# 9
s=base(9,'Week 1 demo readiness','Deliver a scoped, reproducible local proof — not a full retailer platform')
for i,(a,b) in enumerate([('READY','SQLite seed + local API / browser'),('READY','6 deterministic tests passing'),('READY','Four demo question flows'),('TO VERIFY','LangGraph + Chroma full-mode run'),('TO VERIFY','RAGAS scores + LangSmith trace')]):
 y=2.34+i*.74;shape(s,.88,y,11.5,.59,CARD);txt(s,a,1.12,y+.12,2,.3,13,TEAL if a=='READY' else GOLD,True);txt(s,b,3.2,y+.1,8.8,.36,15,TEXT)
shape(s,.87,6.37,11.55,.43,'#204A55');txt(s,'DEMO GOAL: accurate stock + revenue + bestsellers + sourced product answer.',1.1,6.43,11,.3,12,TEXT,True)
path=Path(__file__).resolve().parents[1]/'slides'/'RetailIQ_Flezi_Week1_Demo.pptx';path.parent.mkdir(exist_ok=True);R.save(path)
print(path,len(R.slides),'slides')
